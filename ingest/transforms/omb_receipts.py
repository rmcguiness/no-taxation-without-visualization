"""Transform OMB Historical Table 2.1 (Receipts by Source) to tidy rows.

Layout (sheet 'Table'):
  row 1  title, row 2 '(in millions of dollars)'
  row 3  column groups: Fiscal Year | Individual Income Taxes |
         Corporation Income Taxes | Social Insurance and Retirement Receipts |
         Excise Taxes | Other | Total Receipts
  row 4  sub-headers (Total / On-Budget / Off-Budget) for grouped columns
  rows 5+  one row per fiscal year ('1934' … '2031', 'TQ', '2026 estimate' …)
  footer rows starting with '(' are footnotes.

Phase 1 keeps the Total column of grouped sources (on/off-budget splits are
not ingested). The TQ row is excluded (see transforms.common.parse_fiscal_year).
"""

from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

import openpyxl

from transforms.common import SeriesRow, ValidationReport, cell_to_usd, parse_fiscal_year

TOTAL_CODE = "TOTAL"

# (category_code, header text on row 3). Codes are our stable slugs — OMB
# assigns no codes in Table 2.1. Order matters: components precede the total.
COLUMNS: list[tuple[str, str]] = [
    ("individual_income_taxes", "Individual Income Taxes"),
    ("corporation_income_taxes", "Corporation Income Taxes"),
    ("social_insurance_retirement", "Social Insurance and Retirement Receipts"),
    ("excise_taxes", "Excise Taxes"),
    ("other", "Other"),
    (TOTAL_CODE, "Total Receipts"),
]

_FOOTNOTE_MARKER = re.compile(r"\s*\(\d+\)\s*$")


def _clean_header(text: str) -> str:
    """Strip footnote markers: 'Corporation Income Taxes (1)' → '…Taxes'."""
    return _FOOTNOTE_MARKER.sub("", text.strip())


def _find_header_row(ws) -> int:
    for r in range(1, 11):
        if str(ws.cell(r, 1).value or "").strip() == "Fiscal Year":
            return r
    raise ValueError("Table 2.1 layout changed: no 'Fiscal Year' header row found")


def _map_columns(ws, header_row: int) -> dict[str, int]:
    headers = {
        _clean_header(str(ws.cell(header_row, c).value)): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(header_row, c).value is not None
    }
    mapping: dict[str, int] = {}
    for code, header in COLUMNS:
        if header not in headers:
            raise ValueError(f"Table 2.1 layout changed: column {header!r} not found")
        mapping[code] = headers[header]
    return mapping


def parse_receipts(path: Path) -> list[SeriesRow]:
    ws = openpyxl.load_workbook(path, data_only=True)["Table"]
    header_row = _find_header_row(ws)
    columns = _map_columns(ws, header_row)
    names = {code: header for code, header in COLUMNS}

    rows: list[SeriesRow] = []
    for r in range(header_row + 2, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label is None:
            continue
        parsed = parse_fiscal_year(label)
        if parsed is None:
            continue  # TQ row or footnote
        fiscal_year, is_projection = parsed
        for code, col in columns.items():
            amount = cell_to_usd(ws.cell(r, col).value)
            if amount is None:
                continue
            rows.append(
                SeriesRow(
                    government_level="federal",
                    metric="receipt",
                    category_code=code,
                    category_name=names[code],
                    parent_category_code=None if code == TOTAL_CODE else TOTAL_CODE,
                    fiscal_year=fiscal_year,
                    amount=amount,
                    is_projection=is_projection,
                )
            )
    if not rows:
        raise ValueError(f"no receipt rows parsed from {path}")
    return rows


def validate_receipts(rows: list[SeriesRow]) -> ValidationReport:
    """Every year's component sum must match the published Total Receipts."""
    report = ValidationReport()
    by_year: dict[int, dict[str, Decimal]] = {}
    for row in rows:
        by_year.setdefault(row.fiscal_year, {})[row.category_code] = row.amount
    for year, cats in sorted(by_year.items()):
        if TOTAL_CODE not in cats:
            report.failures.append(f"FY{year}: no Total Receipts value")
            continue
        component_sum = sum(
            (v for code, v in cats.items() if code != TOTAL_CODE), Decimal(0)
        )
        report.check(f"receipts FY{year}", cats[TOTAL_CODE], component_sum)
    return report
