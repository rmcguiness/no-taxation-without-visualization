"""Transform OMB Historical Table 3.2 (Outlays by Function and Subfunction).

Layout (sheet 'Table'): years are columns ('1962' … '2031', a 'TQ' column,
'2026 estimate' …); rows are grouped by budget function:

  '050 National Defense:'            function header (no values)
  '051 Department of Defense-Military:'  subfunction with uncoded detail rows
  'Military Personnel' …                 detail rows (skipped in Phase 1)
  '051 Subtotal, Department of …'        subfunction value row
  '053 Atomic energy defense activities' plain subfunction value row
  'Total, National Defense'              function total row
  '(On-budget)' / '(Off-budget)'         splits (skipped in Phase 1)
  'Total outlays'                        grand total (category_code 'TOTAL')

Two functions (570 Medicare, 650 Social Security) have a single subfunction
and no 'Total,' row in the source; their function-level rows are derived from
that subfunction and marked in `notes`.

Validation: per function, subfunction sums must match the function total; per
year, function totals must match 'Total outlays' — both within the tolerance
documented in transforms.common. Missing cells (dots) count as zero in sums.
"""

from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

import openpyxl

from transforms.common import SeriesRow, ValidationReport, cell_to_usd, parse_fiscal_year

TOTAL_CODE = "TOTAL"

_FUNCTION_HEADER = re.compile(r"^(\d{3}) (.+):$")
_SUBTOTAL = re.compile(r"^(\d{3}) Subtotal, (.+)$")
_SUBFUNCTION = re.compile(r"^(\d{3}) (.+)$")
_FUNCTION_TOTAL = re.compile(r"^Total, (.+)$")
_SKIP_LABELS = re.compile(r"^\((On|Off)-budget\)$", re.IGNORECASE)


def _year_columns(ws, header_row: int) -> dict[int, tuple[int, bool]]:
    """Map column index → (fiscal_year, is_projection); TQ column excluded."""
    years: dict[int, tuple[int, bool]] = {}
    for c in range(2, ws.max_column + 1):
        value = ws.cell(header_row, c).value
        if value is None:
            continue
        parsed = parse_fiscal_year(value)
        if parsed is not None:
            years[c] = parsed
    if not years:
        raise ValueError("Table 3.2 layout changed: no year columns found")
    return years


def _find_header_row(ws) -> int:
    for r in range(1, 11):
        if str(ws.cell(r, 1).value or "").strip() == "Function and Subfunction":
            return r
    raise ValueError("Table 3.2 layout changed: no 'Function and Subfunction' row")


def _emit(rows: list[SeriesRow], ws, r: int, years: dict[int, tuple[int, bool]],
          code: str, name: str, parent: str | None, notes: str | None = None) -> None:
    for col, (fiscal_year, is_projection) in years.items():
        amount = cell_to_usd(ws.cell(r, col).value)
        if amount is None:
            continue
        rows.append(
            SeriesRow(
                government_level="federal",
                metric="outlay",
                category_code=code,
                category_name=name,
                parent_category_code=parent,
                fiscal_year=fiscal_year,
                amount=amount,
                is_projection=is_projection,
                notes=notes,
            )
        )


def parse_outlays(path: Path) -> list[SeriesRow]:
    ws = openpyxl.load_workbook(path, data_only=True)["Table"]
    header_row = _find_header_row(ws)
    years = _year_columns(ws, header_row)

    rows: list[SeriesRow] = []
    current_function: tuple[str, str] | None = None  # (code, name)
    functions_with_total: set[str] = set()
    function_sub_rows: dict[str, list[int]] = {}  # function code → sheet rows of subfunctions

    for r in range(header_row + 1, ws.max_row + 1):
        raw_label = ws.cell(r, 1).value
        if raw_label is None:
            continue
        label = str(raw_label).strip()

        if _SKIP_LABELS.match(label):
            continue

        m = _FUNCTION_HEADER.match(label)
        if m:
            # Function codes end in 0 (050, 150, … 950). A colon-terminated
            # header whose code doesn't (e.g. '051 Department of
            # Defense-Military:') opens a subfunction detail block; its value
            # arrives later on the 'NNN Subtotal, …' row, so skip it here.
            if m.group(1).endswith("0"):
                current_function = (m.group(1), m.group(2))
                function_sub_rows[m.group(1)] = []
            continue

        m = _FUNCTION_TOTAL.match(label)
        if m:
            if current_function is None or m.group(1) != current_function[1]:
                raise ValueError(
                    f"row {r}: total row {label!r} does not match open function "
                    f"{current_function!r}"
                )
            _emit(rows, ws, r, years, current_function[0], current_function[1], None)
            functions_with_total.add(current_function[0])
            continue

        if label == "Total outlays":
            _emit(rows, ws, r, years, TOTAL_CODE, "Total outlays", None)
            current_function = None
            continue

        m = _SUBTOTAL.match(label) or _SUBFUNCTION.match(label)
        if m and current_function is not None:
            sub_code, sub_name = m.group(1), m.group(2)
            _emit(rows, ws, r, years, sub_code, sub_name, current_function[0])
            function_sub_rows[current_function[0]].append(r)
            continue

        # Uncoded detail rows (e.g. 'Military Personnel') and footer notes.
        continue

    # Functions with no 'Total,' row (570 Medicare, 650 Social Security):
    # derive the function series from their single subfunction.
    for code, sub_row_indexes in function_sub_rows.items():
        if code in functions_with_total:
            continue
        if len(sub_row_indexes) != 1:
            raise ValueError(
                f"function {code} has no total row and {len(sub_row_indexes)} "
                "subfunctions; cannot derive its series"
            )
        function_name = _function_name_for(ws, code, header_row)
        _emit(
            rows, ws, sub_row_indexes[0], years, code, function_name, None,
            notes="derived from single subfunction; source has no 'Total,' row",
        )

    if not rows:
        raise ValueError(f"no outlay rows parsed from {path}")
    return rows


def _function_name_for(ws, code: str, header_row: int) -> str:
    for r in range(header_row + 1, ws.max_row + 1):
        label = str(ws.cell(r, 1).value or "").strip()
        m = _FUNCTION_HEADER.match(label)
        if m and m.group(1) == code:
            return m.group(2)
    raise ValueError(f"function header for code {code} not found")


def validate_outlays(rows: list[SeriesRow]) -> ValidationReport:
    """Function totals vs subfunction sums, and grand total vs function sums."""
    report = ValidationReport()

    functions: dict[tuple[str, int], Decimal] = {}
    sub_sums: dict[tuple[str, int], Decimal] = {}
    totals: dict[int, Decimal] = {}
    derived: set[str] = set()

    for row in rows:
        key = (row.category_code, row.fiscal_year)
        if row.category_code == TOTAL_CODE:
            totals[row.fiscal_year] = row.amount
        elif row.parent_category_code is None:
            functions[key] = row.amount
            if row.notes and "derived" in row.notes:
                derived.add(row.category_code)
        else:
            parent_key = (row.parent_category_code, row.fiscal_year)
            sub_sums[parent_key] = sub_sums.get(parent_key, Decimal(0)) + row.amount

    for (code, year), total in sorted(functions.items()):
        if code in derived:
            continue  # function row IS the single subfunction; nothing to check
        report.check(
            f"function {code} FY{year}", total, sub_sums.get((code, year), Decimal(0))
        )

    function_year_sums: dict[int, Decimal] = {}
    for (code, year), amount in functions.items():
        function_year_sums[year] = function_year_sums.get(year, Decimal(0)) + amount
    for year, total in sorted(totals.items()):
        report.check(
            f"total outlays FY{year}", total, function_year_sums.get(year, Decimal(0))
        )
    if not totals:
        report.failures.append("no 'Total outlays' rows parsed")
    return report
